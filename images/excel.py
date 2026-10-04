import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference


top10 = pd.read_csv("q1_top10.csv")
india_monthly = pd.read_csv("q2_india_monthly.csv")
continent = pd.read_csv("q5_continent.csv")


arima = pd.read_csv("results_arima.csv")
ml = pd.read_csv("results_ml.csv")

model_results = pd.concat(
    [arima, ml],
    ignore_index=True
)

out = "covid_report.xlsx"

with pd.ExcelWriter(out, engine="openpyxl") as writer:

    top10.to_excel(
        writer,
        sheet_name="Top10",
        index=False
    )

    india_monthly.to_excel(
        writer,
        sheet_name="India_Monthly",
        index=False
    )

    continent.to_excel(
        writer,
        sheet_name="Continent",
        index=False
    )

    model_results.to_excel(
        writer,
        sheet_name="Model_Results",
        index=False
    )

wb = load_workbook(out)

ws = wb["Top10"]

chart = BarChart()
chart.title = "Top 10 Countries by Total Cases"
chart.y_axis.title = "Total Cases"
chart.x_axis.title = "Country"

data = Reference(
    ws,
    min_col=2,
    min_row=1,
    max_row=11
)

categories = Reference(
    ws,
    min_col=1,
    min_row=2,
    max_row=11
)

chart.add_data(
    data,
    titles_from_data=True
)

chart.set_categories(categories)

ws.add_chart(chart, "D2")

wb.save(out)

print("Excel report created:", out)